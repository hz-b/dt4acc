#!/usr/bin/env python3
"""
Convert EPICS to Tango presentation from Markdown to PDF and PowerPoint formats.
Requires: pip install markdown pdfkit python-pptx
"""

import os
import sys
import markdown
import pdfkit
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import re

def markdown_to_html(markdown_file):
    """Convert Markdown to HTML"""
    with open(markdown_file, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Configure markdown with extensions
    md = markdown.Markdown(extensions=['toc', 'tables', 'fenced_code', 'codehilite'])
    html_content = md.convert(md_content)
    
    # Create full HTML document
    full_html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <title>EPICS to Tango Migration Presentation</title>
        <style>
            body {{
                font-family: Arial, sans-serif;
                line-height: 1.6;
                max-width: 1200px;
                margin: 0 auto;
                padding: 20px;
            }}
            h1 {{
                color: #2c3e50;
                border-bottom: 3px solid #3498db;
                padding-bottom: 10px;
            }}
            h2 {{
                color: #34495e;
                border-bottom: 2px solid #ecf0f1;
                padding-bottom: 5px;
            }}
            h3 {{
                color: #7f8c8d;
            }}
            code {{
                background-color: #f8f9fa;
                padding: 2px 4px;
                border-radius: 3px;
                font-family: 'Courier New', monospace;
            }}
            pre {{
                background-color: #f8f9fa;
                padding: 15px;
                border-radius: 5px;
                border-left: 4px solid #3498db;
                overflow-x: auto;
            }}
            table {{
                border-collapse: collapse;
                width: 100%;
                margin: 20px 0;
            }}
            th, td {{
                border: 1px solid #ddd;
                padding: 12px;
                text-align: left;
            }}
            th {{
                background-color: #3498db;
                color: white;
            }}
            .toc {{
                background-color: #ecf0f1;
                padding: 20px;
                border-radius: 5px;
                margin: 20px 0;
            }}
        </style>
    </head>
    <body>
        {html_content}
    </body>
    </html>
    """
    
    return full_html

def html_to_pdf(html_content, output_file):
    """Convert HTML to PDF"""
    try:
        options = {
            'page-size': 'A4',
            'margin-top': '0.75in',
            'margin-right': '0.75in',
            'margin-bottom': '0.75in',
            'margin-left': '0.75in',
            'encoding': "UTF-8",
            'no-outline': None,
            'enable-local-file-access': None
        }
        pdfkit.from_string(html_content, output_file, options=options)
        print(f"✅ PDF created: {output_file}")
    except Exception as e:
        print(f"❌ PDF creation failed: {e}")
        print("💡 Install wkhtmltopdf: https://wkhtmltopdf.org/downloads.html")

def parse_markdown_sections(markdown_file):
    """Parse Markdown file into sections for PowerPoint"""
    with open(markdown_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    sections = []
    current_section = None
    current_content = []
    
    lines = content.split('\n')
    
    for line in lines:
        if line.startswith('# '):
            if current_section:
                sections.append((current_section, '\n'.join(current_content)))
            current_section = line[2:].strip()
            current_content = []
        elif line.startswith('## '):
            if current_section:
                sections.append((current_section, '\n'.join(current_content)))
            current_section = line[3:].strip()
            current_content = []
        elif line.startswith('### '):
            if current_section:
                sections.append((current_section, '\n'.join(current_content)))
            current_section = line[4:].strip()
            current_content = []
        else:
            current_content.append(line)
    
    if current_section:
        sections.append((current_section, '\n'.join(current_content)))
    
    return sections

def create_powerpoint(sections, output_file):
    """Create PowerPoint presentation from sections"""
    prs = Presentation()
    
    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "EPICS to Tango Migration"
    subtitle.text = "Complete Technical Presentation\n\nBESSY II Accelerator Control System\nAI-Powered Tango Agent"
    
    # Content slides
    for section_title, content in sections:
        if not section_title or section_title in ['Table of Contents', 'Executive Summary']:
            continue
            
        # Create new slide
        bullet_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(bullet_slide_layout)
        
        # Add title
        title = slide.shapes.title
        title.text = section_title
        
        # Add content
        content_placeholder = slide.placeholders[1]
        tf = content_placeholder.text_frame
        tf.clear()
        
        # Process content
        content_lines = content.split('\n')
        for line in content_lines:
            line = line.strip()
            if not line or line.startswith('---'):
                continue
                
            # Clean up markdown formatting
            line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)  # Remove bold
            line = re.sub(r'\*(.*?)\*', r'\1', line)      # Remove italic
            line = re.sub(r'`(.*?)`', r'\1', line)        # Remove code formatting
            line = re.sub(r'^#+\s*', '', line)            # Remove headers
            line = re.sub(r'^\s*[-*]\s*', '• ', line)     # Convert lists
            
            if line:
                p = tf.add_paragraph()
                p.text = line
                p.font.size = Pt(14)
                p.font.name = 'Arial'
        
        # Limit content length
        if len(tf.paragraphs) > 10:
            # Keep only first 10 paragraphs
            for i in range(len(tf.paragraphs) - 1, 9, -1):
                tf._element.remove(tf.paragraphs[i]._element)
    
    # Save presentation
    prs.save(output_file)
    print(f"✅ PowerPoint created: {output_file}")

def main():
    """Main conversion function"""
    markdown_file = "EPICS_TO_TANGO_PRESENTATION.md"
    
    if not os.path.exists(markdown_file):
        print(f"❌ Markdown file not found: {markdown_file}")
        return
    
    print("🔄 Converting Markdown to HTML...")
    html_content = markdown_to_html(markdown_file)
    
    print("🔄 Converting HTML to PDF...")
    html_to_pdf(html_content, "EPICS_TO_TANGO_PRESENTATION.pdf")
    
    print("🔄 Converting to PowerPoint...")
    sections = parse_markdown_sections(markdown_file)
    create_powerpoint(sections, "EPICS_TO_TANGO_PRESENTATION.pptx")
    
    print("\n🎉 Conversion complete!")
    print("📄 Files created:")
    print("   - EPICS_TO_TANGO_PRESENTATION.pdf")
    print("   - EPICS_TO_TANGO_PRESENTATION.pptx")
    print("\n💡 Note: For PDF conversion, you may need to install wkhtmltopdf")
    print("   Download from: https://wkhtmltopdf.org/downloads.html")

if __name__ == "__main__":
    main()
